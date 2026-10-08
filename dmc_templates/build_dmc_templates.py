# -*- coding: utf-8 -*-
"""生成地接社（DMC）资料与报价填写模板：韩国版（中/英/韩表头，按 korea/kb 预填）+ 台湾版（繁中/英表头，按 seed_taiwan 预填）。
用法: python3 build_dmc_templates.py [--kb ../korea/kb] [--out .]
依赖: openpyxl；台湾版繁体转换用 opencc-python-reimplemented（pip install opencc-python-reimplemented）。
字段对齐 docs/02_lark_product_base_schema.md 的 Suppliers / Quotes / Quote_Lines，以及 korea/kb 的 hotels / attractions / city_graph / seasons，
地接填回后 A3 报价转录官可按第 4 行的字段 key 直接读入。
"""
import os, sys, json, argparse, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import seed_taiwan as TW

try:
    from opencc import OpenCC
    _cc = OpenCC("s2twp").convert
    _TW_FIX = [("臺", "台"), ("酒店", "飯店"), ("出租車", "計程車"), ("大巴", "遊覽車")]
    def _T(s):
        s = _cc(s)
        for a, b in _TW_FIX: s = s.replace(a, b)
        return s
except Exception:   # 没装 opencc 时台湾版保留简体
    _T = lambda s: s

VERSION = "v1.0 · 2026-10-08"
FONT = "Arial"
F_TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
F_NOTE = Font(name=FONT, size=9, italic=True, color="595959")
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_KEY = Font(name=FONT, size=8, color="808080")
F_BODY = Font(name=FONT, size=10)
F_EX = Font(name=FONT, size=10, italic=True, color="7F7F7F")
F_PRE = Font(name=FONT, size=10, color="000000")
F_FORMULA = Font(name=FONT, size=10, color="000000")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")     # 黄：地接填写
FILL_PRE = PatternFill("solid", fgColor="EDEDED")       # 灰：WEBUY 预填，勿改
FILL_EX = PatternFill("solid", fgColor="F8F8F8")        # 示例行
FILL_CALC = PatternFill("solid", fgColor="E2EFDA")      # 绿：公式自动计算
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

HEAD_ROW, KEY_ROW, EX_ROW, FIRST = 3, 4, 5, 6   # 第 3 行表头、第 4 行字段 key、第 5 行示例、第 6 行起填写


class Ctx:
    def __init__(self, market):
        self.market = market                      # "KR" / "TW"
        self.kr = market == "KR"
        self.cur = "KRW" if self.kr else "TWD"

    def t(self, s):                               # 台湾版转繁体
        return s if self.kr else _T(s)

    def head(self, cn, en, ko):
        cn = self.t(cn)
        return f"{cn}\n{en}" + (f"\n{ko}" if self.kr and ko else "")


def col(key, cn, en, ko="", w=14, kind="input", dv=None, note=None, num=None):
    """kind: pre（WEBUY 预填）/ input（地接填写）/ calc（公式）。dv: 下拉选项名（Lists 表）。"""
    return dict(key=key, cn=cn, en=en, ko=ko, w=w, kind=kind, dv=dv, note=note, num=num)


def sheet(wb, ctx, name, title, note, cols, example, pre_rows=(), blank=30, formulas=None, lists=None):
    ws = wb.create_sheet(ctx.t(name))
    n = len(cols)
    ws["A1"] = ctx.t(title); ws["A1"].font = F_TITLE
    ws["A2"] = ctx.t(note); ws["A2"].font = F_NOTE; ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(n, 6)); ws.row_dimensions[2].height = 42
    for j, c in enumerate(cols, 1):
        h = ws.cell(HEAD_ROW, j, ctx.head(c["cn"], c["en"], c["ko"])); h.font = F_HEAD; h.fill = FILL_HEAD; h.alignment = CENTER; h.border = BORDER
        if c.get("note"): h.comment = Comment(ctx.t(c["note"]), "WEBUY")
        k = ws.cell(KEY_ROW, j, c["key"]); k.font = F_KEY; k.alignment = CENTER
        ws.column_dimensions[get_column_letter(j)].width = c["w"]
    ws.row_dimensions[HEAD_ROW].height = 58 if ctx.kr else 44
    # 示例行
    for j, c in enumerate(cols, 1):
        v = example.get(c["key"], "")
        cell = ws.cell(EX_ROW, j, ctx.t(v) if isinstance(v, str) else v)
        cell.font = F_EX; cell.fill = FILL_EX; cell.border = BORDER; cell.alignment = WRAP
        if c.get("num"): cell.number_format = c["num"]
    ws.cell(EX_ROW, 1).comment = Comment(ctx.t("示例行（Example）：展示填写格式，不会被读入，可保留。"), "WEBUY")
    rows = list(pre_rows) + [{}] * blank
    for i, r in enumerate(rows):
        rr = FIRST + i
        for j, c in enumerate(cols, 1):
            cell = ws.cell(rr, j); cell.border = BORDER; cell.alignment = WRAP; cell.font = F_BODY
            if c["kind"] == "calc" and formulas and c["key"] in formulas:
                cell.value = formulas[c["key"]](rr); cell.fill = FILL_CALC; cell.font = F_FORMULA
            elif c["kind"] == "pre" and c["key"] in r:
                v = r[c["key"]]; cell.value = ctx.t(v) if isinstance(v, str) else v; cell.fill = FILL_PRE; cell.font = F_PRE
            elif c["kind"] == "pre":
                cell.fill = FILL_INPUT      # 空白新增行：预填列也由地接填
            else:
                v = r.get(c["key"])
                if v is not None: cell.value = ctx.t(v) if isinstance(v, str) else v
                cell.fill = FILL_INPUT
            if c.get("num"): cell.number_format = c["num"]
    last = FIRST + len(rows) - 1
    # 下拉
    for j, c in enumerate(cols, 1):
        if c.get("dv") and lists:
            ref = lists[c["dv"]]
            dv = DataValidation(type="list", formula1=ref, allow_blank=True, showErrorMessage=False)
            ws.add_data_validation(dv); dv.add(f"{get_column_letter(j)}{EX_ROW}:{get_column_letter(j)}{last}")
    ws.freeze_panes = ws.cell(FIRST - 1 if False else EX_ROW, 3)
    ws.auto_filter.ref = f"A{KEY_ROW}:{get_column_letter(n)}{last}"
    ws.sheet_view.zoomScale = 90
    return ws


# ---------------------------------------------------------------- 下拉选项
def build_lists(wb, ctx):
    ws = wb.create_sheet("Lists")
    L = {
        "yn": ["Y", "N"],
        "season": ["淡季 Low", "平季 Shoulder", "旺季 Peak", "节假 Holiday"],
        "currency": [ctx.cur, "USD", "SGD"],
        "star": ["3★", "4★", "5★", "Resort", "Hanok / 韩屋" if ctx.kr else "民宿 B&B", "Glamping"],
        "meal": ["B", "L", "D", "Snack"],
        "suit": ["Group 团体", "FIT 散客", "Both 均可"],
        "tier": ["10–15", "16–20", "21–25", "26–30", "31–35", "36+"],
        "lang": ["中文 Mandarin", "English", "粤语 Cantonese", "Malay", "日本語"],
        "vehicle": ["9 seats (van)", "15 seats", "25 seats", "35 seats", "45 seats"],
        "shoptype": ["购物站 Shopping stop", "自费 Optional tour"],
        "level": ["核心 Core", "常见 Common", "可选 Optional", "地接推荐 DMC pick"],
    }
    if not ctx.kr:
        L = {k: [_T(x) for x in v] for k, v in L.items()}
    refs = {}
    ws["A1"] = ctx.t("下拉选项（勿改）"); ws["A1"].font = F_TITLE
    for j, (k, vals) in enumerate(L.items(), 1):
        ws.cell(3, j, k).font = F_HEAD; ws.cell(3, j).fill = FILL_HEAD
        for i, v in enumerate(vals, 4): ws.cell(i, j, v).font = F_BODY
        c = get_column_letter(j); refs[k] = f"Lists!${c}$4:${c}${3 + len(vals)}"
        ws.column_dimensions[c].width = 18
    # 参考汇率（报价表的 SGD 参考价公式引用这里）
    r0 = 20
    ws.cell(r0, 1, ctx.t("参考汇率（1 单位外币 = ? SGD）")).font = Font(name=FONT, bold=True)
    fx = [(ctx.cur, 0.00094 if ctx.kr else 0.041), ("USD", 1.29), ("SGD", 1.0)]
    for i, (c, v) in enumerate(fx, r0 + 1):
        ws.cell(i, 1, c).font = F_BODY
        x = ws.cell(i, 2, v); x.font = Font(name=FONT, color="0000FF"); x.fill = PatternFill("solid", fgColor="FFFF00"); x.number_format = "0.00000"
    ws.cell(r0 + 4, 1, ctx.t("假设：2026-10 参考汇率，仅用于报价表「SGD 参考价」列；正式成本以 Lark FX_Rates 表为准（A4 自动重算）。")).font = F_NOTE
    refs["fx_range"] = (f"Lists!$A${r0 + 1}:$A${r0 + 3}", f"Lists!$B${r0 + 1}:$B${r0 + 3}")
    ws.sheet_state = "visible"
    return refs


# ---------------------------------------------------------------- 各表
def build(ctx, kb, out_path):
    wb = Workbook(); wb.remove(wb.active)
    readme = wb.create_sheet(ctx.t("填写说明"))
    lists = build_lists(wb, ctx)
    cur = ctx.cur
    country = "韩国" if ctx.kr else "台湾"

    # 1 供应商信息（Suppliers + Quotes 头）
    ws = wb.create_sheet(ctx.t("1 供应商信息"))
    ws["A1"] = ctx.t(f"1 供应商信息 Supplier Profile（{country}地接）"); ws["A1"].font = F_TITLE
    ws["A2"] = ctx.t("对应 Lark Suppliers 表 + Quotes 报价单表头。黄色格子请填写；一家地接填一份。"); ws["A2"].font = F_NOTE
    items = [
        ("supplier_name", "公司名称", "Company name", "회사명", "ABC Travel Co., Ltd."),
        ("license_no", "旅行社执照号", "Travel agency license no.", "여행업 등록번호", "제2026-000001호" if ctx.kr else "交觀甲 0000"),
        ("address", "公司地址", "Address", "주소", ""),
        ("contact_person", "业务联系人 / 职位", "Sales contact / title", "담당자 / 직책", "Kim Min-ji / Sales Manager" if ctx.kr else "王小明 / 業務經理"),
        ("contact_channels", "WhatsApp / 微信 / " + ("KakaoTalk" if ctx.kr else "LINE") + " / Email", "WhatsApp / WeChat / " + ("KakaoTalk" if ctx.kr else "LINE") + " / Email", "연락처", ""),
        ("emergency_24h", "24 小时紧急联系电话", "24h emergency number", "24시간 비상연락처", ""),
        ("regions_covered", "可操作区域", "Regions covered", "운영 가능 지역", "全韩国 + 济州" if ctx.kr else "全台灣本島 + 離島（澎湖/金門）"),
        ("guide_languages", "导游语种与人数", "Guide languages & headcount", "가이드 언어 / 인원", "中文 20 / 英文 8 / 马来 2"),
        ("own_coaches", "自有车队（车型 × 数量）", "Own fleet (type × qty)", "보유 차량", "45 座 × 10；25 座 × 4"),
        ("sg_my_experience", "接待新加坡/马来西亚团经验（年 / 年团量）", "SG/MY market experience", "싱가포르·말레이시아 단체 경험", "8 年 / 约 300 团"),
        ("halal_capable", "可否操作清真 / 素食团", "Halal / vegetarian capability", "할랄 / 채식 가능 여부", "清真餐厅 6 家可用；素食可安排"),
        ("quote_currency", "报价币种", "Quote currency", "견적 통화", cur),
        ("quote_valid_from", "报价有效期 起", "Valid from", "유효기간 시작", "2027-01-01"),
        ("quote_valid_until", "报价有效期 止", "Valid until", "유효기간 종료", "2027-12-31"),
        ("payment_terms", "付款条件（订金 / 尾款期限）", "Payment terms", "결제 조건", "出发前 30 天付 30% 订金，出发前 7 天付清"),
        ("cancellation_policy", "取消政策", "Cancellation policy", "취소 규정", "出发前 14 天取消收 30%……"),
        ("liability_insurance", "旅行社责任险（保额）", "Liability insurance (coverage)", "여행자 배상책임보험", ""),
        ("price_includes_tax", "报价是否含税（VAT）", "Price incl. VAT?", "부가세 포함 여부", "Y"),
        ("notes", "其他说明", "Other notes", "기타", ""),
    ]
    for j, (h, w) in enumerate([("字段 Field", 34), ("填写 Value", 46), ("示例 Example", 40), ("key", 20)], 1):
        c = ws.cell(3, j, ctx.t(h)); c.font = F_HEAD; c.fill = FILL_HEAD; c.alignment = CENTER; ws.column_dimensions[get_column_letter(j)].width = w
    for i, (k, cn, en, ko, ex) in enumerate(items, 4):
        a = ws.cell(i, 1, ctx.head(cn, en, ko).replace("\n", " / ")); a.font = Font(name=FONT, bold=True, size=10); a.alignment = WRAP; a.border = BORDER
        b = ws.cell(i, 2); b.fill = FILL_INPUT; b.border = BORDER; b.alignment = WRAP
        c = ws.cell(i, 3, ctx.t(ex)); c.font = F_EX; c.alignment = WRAP; c.border = BORDER
        d = ws.cell(i, 4, k); d.font = F_KEY
    dv = DataValidation(type="list", formula1=lists["currency"], allow_blank=True); ws.add_data_validation(dv)
    dv.add(f"B{4 + [x[0] for x in items].index('quote_currency')}")
    ws.freeze_panes = "B4"

    # 2 酒店
    hotel_cols = [
        col("city", "城市", "City", "도시", 12, "pre"), col("hotel_name_en", "酒店名称（英文）", "Hotel name (EN)", "호텔명 (영문)", 30, "pre"),
        col("webuy_tier", "WEBUY 档次参考", "Tier seen in market", "시장 등급", 14, "pre", note="竞品线路里该酒店出现的档次：group_standard=新加坡跟团标准，premium=小团高端，fit=自由行，join_in=地接拼团。"),
        col("hotel_name_local", "当地名称", "Local name", "현지명" if ctx.kr else "", 18), col("star", "星级", "Star", "등급", 9, dv="star"),
        col("area", "区域 / 商圈", "Area", "지역", 14), col("room_type", "房型", "Room type", "객실 타입", 14),
        col("season", "季节档", "Season", "시즌", 12, dv="season"), col("date_from", "适用日期 起", "From", "적용 시작", 11, num="yyyy-mm-dd"),
        col("date_to", "适用日期 止", "To", "적용 종료", 11, num="yyyy-mm-dd"),
        col("twin_per_room", f"双人房 / 间 / 晚（{cur}）", f"Twin per room/night ({cur})", "트윈 1실 1박", 13, num="#,##0"),
        col("single_per_room", f"单人房 / 间 / 晚（{cur}）", f"Single per room/night ({cur})", "싱글 1실 1박", 13, num="#,##0"),
        col("triple_per_room", f"三人房 / 间 / 晚（{cur}）", f"Triple per room/night ({cur})", "트리플 1실 1박", 13, num="#,##0"),
        col("extra_bed", f"加床（{cur}）", f"Extra bed ({cur})", "엑스트라 베드", 11, num="#,##0"),
        col("breakfast_incl", "含早餐", "Breakfast incl.", "조식 포함", 9, dv="yn"),
        col("breakfast_price", f"早餐单价（{cur}）", f"Breakfast / pax ({cur})", "조식 단가", 11, num="#,##0"),
        col("child_policy", "儿童政策（年龄 / 不占床）", "Child policy", "아동 규정", 20),
        col("group_min_rooms", "团体价最少间数", "Min rooms for group rate", "단체 최소 객실", 11),
        col("coach_parking", "大巴可停", "Coach parking", "버스 주차", 9, dv="yn"),
        col("allotment", "可保留房量（间 / 晚）", "Allotment (rooms/night)", "블록 객실 수", 12),
        col("release_days", "释放期（出发前天数）", "Release (days before)", "릴리즈 기한", 11),
        col("notes", "备注（加价日 / 施工 / 新开）", "Notes", "비고", 26),
    ]
    if ctx.kr:
        hotels = json.load(open(os.path.join(kb, "hotels.json"), encoding="utf-8"))
        pre = []
        for h in hotels:
            for tier, names in h["tiers"].items():
                for nme in names:
                    pre.append({"city": f"{h['city_cn']} {h['city_key'].title()}", "hotel_name_en": nme, "webuy_tier": tier})
        hex_ = {"city": "首尔 Seoul", "hotel_name_en": "Novotel Ambassador Seoul Yongsan", "webuy_tier": "group_standard", "hotel_name_local": "노보텔 앰배서더 서울 용산",
                "star": "4★", "area": "龙山 Yongsan", "room_type": "Superior Twin", "season": "平季 Shoulder", "date_from": datetime.date(2027, 3, 1), "date_to": datetime.date(2027, 6, 30),
                "twin_per_room": 180000, "single_per_room": 170000, "triple_per_room": 240000, "extra_bed": 50000, "breakfast_incl": "Y", "breakfast_price": 35000,
                "child_policy": "12 岁以下不占床免费（不含早）", "group_min_rooms": 8, "coach_parking": "Y", "allotment": 20, "release_days": 21, "notes": "4/1–4/10 樱花期加价 20%"}
        hnote = f"已预填竞品在售线路出现过的 {len(pre)} 家酒店（灰色）。请对能操作的酒店填价，不能操作的留空；可在下方黄色空行补充推荐酒店。价格为每间每晚净价。"
    else:
        pre = []
        hex_ = {"city": "台北 Taipei", "hotel_name_en": "Caesar Park Hotel Taipei", "webuy_tier": "", "hotel_name_local": "台北凱撒大飯店", "star": "4★", "area": "台北車站",
                "room_type": "Superior Twin", "season": "平季 Shoulder", "date_from": datetime.date(2027, 3, 1), "date_to": datetime.date(2027, 6, 30), "twin_per_room": 4200,
                "single_per_room": 4000, "triple_per_room": 5200, "extra_bed": 1000, "breakfast_incl": "Y", "breakfast_price": 600, "child_policy": "12 歲以下不佔床免費（不含早）",
                "group_min_rooms": 8, "coach_parking": "Y", "allotment": 15, "release_days": 21, "notes": "跨年、國慶加價"}
        hnote = "请按城市填写可操作的团体酒店（每家每季一行）；建议覆盖台北、宜兰/礁溪、花莲、台东、垦丁、高雄、台南、嘉义/阿里山、日月潭、台中、清境。价格为每间每晚净价。"
    sheet(wb, ctx, "2 酒店", "2 酒店 Hotels", hnote, hotel_cols, hex_, pre, blank=40, lists=lists)

    # 3 景点门票
    att_cols = [
        col("city", "城市", "City", "도시", 13, "pre"), col("name_en", "景点（英文）", "Attraction (EN)", "관광지 (영문)", 30, "pre"),
        col("name_cn", "景点（中文）", "Attraction (CN)", "관광지 (중문)", 22, "pre"), col("category", "类别", "Category", "분류", 12, "pre"),
        col("webuy_level", "市场热度", "Market level", "시장 인기도", 12, "pre", dv="level", note="核心=出现在 ≥4 条竞品线路；常见=2–3 条；可选=1 条。"),
        col("name_local", "当地名称", "Local name", "현지명" if ctx.kr else "", 18), col("duration_h", "建议停留（小时）", "Suggested stay (h)", "권장 체류", 10, num="0.0"),
        col("opening_hours", "开放时间", "Opening hours", "운영 시간", 14), col("closed_days", "休息日", "Closed days", "휴무일", 12),
        col("adult_ticket", f"成人门票（{cur}）", f"Adult ticket ({cur})", "성인 입장료", 11, num="#,##0"),
        col("child_ticket", f"儿童门票（{cur}）", f"Child ticket ({cur})", "아동 입장료", 11, num="#,##0"),
        col("group_rate", f"团体价 / 人（{cur}）", f"Group rate/pax ({cur})", "단체 요금", 11, num="#,##0"),
        col("group_min_pax", "团体价最少人数", "Group min pax", "단체 최소 인원", 10),
        col("free_policy", "免费政策（导游 / 司机 / FOC）", "Free policy", "무료 규정", 16),
        col("reservation_days", "需预约（提前天数，0=不需）", "Booking lead days", "예약 필요 일수", 11),
        col("coach_parking", "大巴可停", "Coach parking", "버스 주차", 9, dv="yn"),
        col("best_months", "最佳月份 / 季节限制", "Best months / seasonal limits", "추천 시기", 16),
        col("suitable", "适合", "Suitable for", "적합 유형", 11, dv="suit"),
        col("halal_friendly", "清真友好（祈祷室 / 餐）", "Muslim-friendly", "무슬림 친화", 10, dv="yn"),
        col("notes", "备注（施工 / 关闭 / 替代景点）", "Notes", "비고", 26),
    ]
    if ctx.kr:
        A = json.load(open(os.path.join(kb, "attractions.json"), encoding="utf-8"))
        lv = {"core": "核心 Core", "common": "常见 Common", "optional": "可选 Optional"}
        keep = [a for a in A if not a.get("generic") and a["category"] not in ("shopstop",) and a["tour_count"] >= 1]
        ro = {"Seoul Capital Area": 0, "Jeju": 1, "Gyeongsangnam-do / Busan / Ulsan": 2, "Gyeongsangbuk-do / Daegu": 3, "Gangwon": 4, "Jeolla": 5, "Chungcheong": 6}
        keep.sort(key=lambda a: ({"core": 0, "common": 1, "optional": 2}[a["level"]], ro.get(a["region"], 9), a["city_key"] or "", -a["tour_count"]))
        pre = [{"city": f"{a['city_cn']} {a['city']}", "name_en": a["name_en"], "name_cn": a["name_cn"], "category": a["category_cn"], "webuy_level": lv[a["level"]]} for a in keep]
        aex = {"city": "首尔 Seoul", "name_en": "Gyeongbokgung Palace", "name_cn": "景福宫", "category": "宫殿", "webuy_level": "核心 Core", "name_local": "경복궁",
               "duration_h": 1.5, "opening_hours": "09:00–18:00", "closed_days": "周二", "adult_ticket": 3000, "child_ticket": 0, "group_rate": 2400, "group_min_pax": 10,
               "free_policy": "穿韩服免费入场", "reservation_days": 0, "coach_parking": "Y", "best_months": "全年", "suitable": "Both 均可", "halal_friendly": "N", "notes": "周二休馆时改德寿宫"}
        anote = f"已预填竞品线路出现过的 {len(pre)} 个景点（按 核心 → 常见 → 可选 排序）。请填写门票与团体价、开放与预约要求；不再营业或不建议的请在备注写替代景点；可在下方空行补充地接推荐景点。"
    else:
        pre = [{"city": c, "name_en": en, "name_cn": cn, "category": cat} for c, en, cn, cat, h in TW.ATTRACTIONS]
        aex = {"city": "南投 Nantou", "name_en": "Sun Moon Lake Ropeway", "name_cn": "日月潭纜車", "category": "纜車", "webuy_level": "核心 Core", "name_local": "日月潭纜車",
               "duration_h": 1.5, "opening_hours": "10:30–16:00", "closed_days": "每月第一個週三", "adult_ticket": 300, "child_ticket": 250, "group_rate": 270, "group_min_pax": 20,
               "free_policy": "導遊司機免費", "reservation_days": 3, "coach_parking": "Y", "best_months": "全年", "suitable": "Both 均可", "halal_friendly": "N", "notes": "強風停駛，改遊湖船"}
        anote = f"已预填 {len(pre)} 个台湾团体常用景点（灰色，WEBUY 整理的常见清单，非竞品数据）。请填写门票/团体价与限制；可在下方空行补充推荐景点。"
    sheet(wb, ctx, "3 景点门票", "3 景点门票 Attractions & Tickets", anote, att_cols, aex, pre, blank=30, lists=lists)

    # 4 餐食
    meal_cols = [
        col("city", "城市", "City", "도시", 13), col("meal_name", "餐名 / 菜单", "Meal / menu", "메뉴", 26), col("restaurant", "餐厅名称", "Restaurant", "식당명", 22),
        col("meal_type", "餐别", "Meal", "식사 구분", 8, dv="meal"), col("cuisine", "菜系 / 特色", "Cuisine / highlight", "음식 종류", 16),
        col("price_adult", f"成人 / 人（{cur}）", f"Adult/pax ({cur})", "성인 1인", 11, num="#,##0"), col("price_child", f"儿童 / 人（{cur}）", f"Child/pax ({cur})", "아동 1인", 11, num="#,##0"),
        col("min_pax", "最少人数", "Min pax", "최소 인원", 9), col("capacity", "可容纳人数", "Seating capacity", "수용 인원", 9),
        col("halal", "清真", "Halal", "할랄", 7, dv="yn"), col("vegetarian", "素食可替换", "Vegetarian option", "채식 대체", 9, dv="yn"),
        col("no_beef", "可不含牛肉", "No-beef option", "소고기 제외", 9, dv="yn"), col("coach_parking", "大巴可停", "Coach parking", "버스 주차", 9, dv="yn"),
        col("notes", "备注", "Notes", "비고", 26),
    ]
    if ctx.kr:
        T = json.load(open(os.path.join(kb, "tours.json"), encoding="utf-8"))
        mpre = [{"city": "首尔 Seoul", "meal_name": "参鸡汤 Ginseng chicken soup"}, {"city": "首尔 Seoul", "meal_name": "韩式烤肉（猪/牛）Korean BBQ"},
                {"city": "首尔 Seoul", "meal_name": "部队锅 Army stew"}, {"city": "首尔 Seoul", "meal_name": "韩定食 Hanjeongsik"},
                {"city": "济州 Jeju", "meal_name": "济州黑猪烤肉 Jeju black pork"}, {"city": "济州 Jeju", "meal_name": "带鱼 / 青花鱼套餐 Hairtail & mackerel set"},
                {"city": "济州 Jeju", "meal_name": "海鲜锅 Seafood hotpot"}, {"city": "釜山 Busan", "meal_name": "猪肉汤饭 Dwaeji-gukbap"},
                {"city": "釜山 Busan", "meal_name": "海鲜 / 生鱼片套餐 Sashimi set"}, {"city": "全州 Jeonju", "meal_name": "全州石锅拌饭 Jeonju bibimbap"},
                {"city": "春川 Chuncheon", "meal_name": "春川辣炒鸡排 Dakgalbi"}, {"city": "安东 Andong", "meal_name": "安东炖鸡 Andong jjimdak"},
                {"city": "江原道 Gangwon", "meal_name": "滑雪场自助 / 韩式套餐 Resort buffet"}, {"city": "首尔 Seoul", "meal_name": "清真韩餐 Halal Korean set"}]
        mex = {"city": "首尔 Seoul", "meal_name": "参鸡汤 + 海鲜煎饼", "restaurant": "土俗村参鸡汤", "meal_type": "L", "cuisine": "韩式 / 米其林指南推荐", "price_adult": 22000,
               "price_child": 15000, "min_pax": 10, "capacity": 120, "halal": "N", "vegetarian": "Y", "no_beef": "Y", "coach_parking": "N", "notes": "大巴需停景福宫停车场，步行 5 分钟"}
        mnote = "灰色为竞品团常见团餐（请对应报价、推荐餐厅）；黄色空行补充其他团餐与特色餐。价格为每人净价，含税。"
    else:
        mpre = [{"city": c, "meal_name": m} for c, m, t in TW.MEALS]
        mex = {"city": "台北 Taipei", "meal_name": "台菜合菜（十菜一湯）", "restaurant": "欣葉台菜", "meal_type": "D", "cuisine": "台菜", "price_adult": 650, "price_child": 450,
               "min_pax": 10, "capacity": 200, "halal": "N", "vegetarian": "Y", "no_beef": "Y", "coach_parking": "Y", "notes": "10 人一桌"}
        mnote = "灰色为常见团餐（请对应报价、推荐餐厅）；黄色空行补充其他团餐与特色餐。价格为每人净价，含税。"
    meal_cols[0]["kind"] = "pre"; meal_cols[1]["kind"] = "pre"
    sheet(wb, ctx, "4 餐食", "4 餐食 Meals", mnote, meal_cols, mex, mpre, blank=30, lists=lists)

    # 5 用车与交通
    tr_cols = [
        col("item", "项目 / 路段", "Item / segment", "항목 / 구간", 32, "pre"), col("mode", "方式", "Mode", "교통수단", 18, "pre"),
        col("market_seen", "竞品线路出现次数", "Seen in market tours", "시장 노출", 10, "pre"),
        col("vehicle", "车型", "Vehicle", "차종", 13, dv="vehicle"), col("duration", "车程 / 航程", "Duration", "소요 시간", 10),
        col("price", f"价格（{cur}）", f"Price ({cur})", "요금", 12, num="#,##0"), col("price_unit", "计价单位", "Price unit", "단위", 14),
        col("included", "已含（油/路桥/停车/司机食宿）", "Included", "포함 사항", 20), col("ot_per_hour", f"超时 / 小时（{cur}）", f"Overtime/h ({cur})", "초과 시간당", 11, num="#,##0"),
        col("notes", "备注", "Notes", "비고", 26),
    ]
    if ctx.kr:
        G = json.load(open(os.path.join(kb, "city_graph.json"), encoding="utf-8"))
        cn = {n["key"]: n["name_cn"] for n in G["nodes"]}
        tpre = [{"item": "大巴日租（10 小时，市内）", "mode": "Coach"}, {"item": "大巴超长日（跨道 ≥ 300 km）", "mode": "Coach"},
                {"item": "仁川机场 ICN 接 / 送", "mode": "Coach"}, {"item": "金浦机场 GMP 接 / 送", "mode": "Coach"}, {"item": "金海机场 PUS 接 / 送", "mode": "Coach"},
                {"item": "济州机场 CJU 接 / 送", "mode": "Coach"}, {"item": "济州环岛大巴日租", "mode": "Coach"},
                {"item": "国内段 GMP–CJU（团体票）", "mode": "Flight"}, {"item": "国内段 CJU–PUS（团体票）", "mode": "Flight"}, {"item": "国内段 GMP–PUS（团体票）", "mode": "Flight"},
                {"item": "KTX 首尔–釜山", "mode": "KTX"}, {"item": "KTX 首尔–江陵", "mode": "KTX"}, {"item": "牛岛渡轮（往返 + 车辆）", "mode": "Ferry"},
                {"item": "南怡岛渡船", "mode": "Ferry"}, {"item": "浦项–郁陵岛高速船", "mode": "Ferry"}]
        for e in G["edges"]:
            if e["tour_count"] >= 2:
                tpre.append({"item": f"{cn.get(e['from']) or e['from']} → {cn.get(e['to']) or e['to']}", "mode": " / ".join(e["modes"]), "market_seen": e["tour_count"]})
        tex = {"item": "大巴日租（10 小时，市内）", "mode": "Coach", "market_seen": "", "vehicle": "45 seats", "duration": "10h", "price": 900000, "price_unit": "每车每天",
               "included": "油、路桥、停车；不含司机住宿", "ot_per_hour": 60000, "notes": "跨道过夜另加司机房 1 间"}
        tnote = "灰色为常用用车项目与竞品线路实际出现过的城市路段（次数=出现的线路数）。请按车型填价；国内段机票/KTX 请填团体票价与出票条件。"
    else:
        tpre = [{"item": "遊覽車日租（10 小時）", "mode": "Coach"}, {"item": "桃園機場 TPE 接 / 送", "mode": "Coach"}, {"item": "環島 8 天包車", "mode": "Coach"}] + \
               [{"item": s, "mode": m, "market_seen": ""} for s, m, d in TW.SEGMENTS]
        tex = {"item": "遊覽車日租（10 小時）", "mode": "Coach", "vehicle": "45 seats", "duration": "10h", "price": 13000, "price_unit": "每車每天",
               "included": "油資、過路費、停車費；司機食宿另計", "ot_per_hour": 800, "notes": "阿里山、清境需換中巴"}
        tnote = "灰色为常用用车项目与环岛常见路段。请按车型填价；高铁/台铁请填团体票价与订位规则；山区路段请注明是否需换中巴。"
    sheet(wb, ctx, "5 用车交通", "5 用车与交通 Transport", tnote, tr_cols, tex, tpre, blank=20, lists=lists)

    # 6 导游领队
    g_cols = [
        col("role", "角色", "Role", "역할", 16, "pre"), col("language", "语种", "Language", "언어", 14, dv="lang"),
        col("fee_per_day", f"日薪（{cur}）", f"Fee/day ({cur})", "일당", 12, num="#,##0"), col("meal_allowance", f"餐补 / 天（{cur}）", f"Meal allowance/day ({cur})", "식대", 12, num="#,##0"),
        col("room_needed", "需另付住宿", "Room needed", "숙박 필요", 9, dv="yn"), col("tips_guideline", f"建议小费 / 人 / 天（{cur}）", f"Tips/pax/day ({cur})", "팁 가이드", 14, num="#,##0"),
        col("foc_policy", "FOC 政策（如 15+1）", "FOC policy", "FOC 규정", 14), col("notes", "备注", "Notes", "비고", 30),
    ]
    gpre = [{"role": "中文导游 Mandarin guide"}, {"role": "英文导游 English guide"}, {"role": "司机兼导游 Driver-guide"}, {"role": "送机 / 接机助理 Airport assistant"}]
    gex = {"role": "中文导游 Mandarin guide", "language": "中文 Mandarin", "fee_per_day": 200000 if ctx.kr else 3000, "meal_allowance": 30000 if ctx.kr else 400,
           "room_needed": "Y", "tips_guideline": 10000 if ctx.kr else 300, "foc_policy": "15+1（第 16 位免费）", "notes": "EU Holidays 韩国线向客人收服务费 KRW 90,000 / 人（8D6N）" if ctx.kr else "小費導遊與司機合計"}
    sheet(wb, ctx, "6 导游领队", "6 导游 / 领队 / 小费 Guides & Tips", "填写导游费用与建议小费（WEBUY 会写进行程「不含」与「服务费」）。", g_cols, gex, gpre, blank=8, lists=lists)

    # 7 购物站与自费
    s_cols = [
        col("name", "名称", "Name", "명칭", 26, "pre"), col("type", "类型", "Type", "유형", 16, "pre", dv="shoptype"), col("city", "城市", "City", "도시", 12),
        col("stay_minutes", "停留（分钟）", "Stay (min)", "체류 (분)", 10), col("price", f"自费价格 / 人（{cur}）", f"Optional price/pax ({cur})", "선택관광 요금", 12, num="#,##0"),
        col("rebate_per_head", f"人头费 / 人（{cur}）", f"Rebate per head ({cur})", "인두비", 12, num="#,##0", note="购物站给团的人头费 / 回佣（WEBUY 内部核算用，不对客）。"),
        col("commission_pct", "佣金比例", "Commission %", "커미션", 10, num="0%"), col("mandatory", "是否必进", "Mandatory", "필수 여부", 9, dv="yn"),
        col("notes", "备注（可否免购物 / 免购物加价）", "Notes (no-shopping surcharge)", "비고", 30),
    ]
    if ctx.kr:
        spre = [{"name": "高丽人参店 Ginseng", "type": "购物站 Shopping stop"}, {"name": "护肝宝店 Healthy liver (Hovenia)", "type": "购物站 Shopping stop"},
                {"name": "化妆品店 Cosmetics", "type": "购物站 Shopping stop"}, {"name": "紫水晶店 Amethyst", "type": "购物站 Shopping stop"},
                {"name": "免税店 Duty-free", "type": "购物站 Shopping stop"}, {"name": "头皮护理 Scalp care", "type": "购物站 Shopping stop"},
                {"name": "乐天超市 Lotte Mart", "type": "购物站 Shopping stop"}, {"name": "乱打秀 NANTA show", "type": "自费 Optional tour"},
                {"name": "汉江游船 Han River cruise", "type": "自费 Optional tour"}, {"name": "乐天世界塔 Seoul Sky", "type": "自费 Optional tour"},
                {"name": "釜山游艇 Busan yacht", "type": "自费 Optional tour"}, {"name": "滑雪装备 + 教练 Ski package", "type": "自费 Optional tour"}]
        sex = {"name": "高丽人参店 Ginseng", "type": "购物站 Shopping stop", "city": "首尔 Seoul", "stay_minutes": 60, "price": "", "rebate_per_head": 20000,
               "commission_pct": 0.1, "mandatory": "Y", "notes": "全程不进购物站每人加 KRW 60,000"}
        snote = f"竞品跟团 {round(100 * json.load(open(os.path.join(kb, 'stats.json')))['shopping_stop_share_group'])}% 含购物站（人参 / 护肝宝 / 化妆品最常见）。请列出你们的标准购物站、人头费与「免购物」加价，以便 WEBUY 做纯玩与常规两个版本。"
    else:
        spre = [{"name": n, "type": "购物站 Shopping stop" if t == "購物站" else "自费 Optional tour"} for n, t in TW.SHOPPING]
        sex = {"name": "茶葉（高山茶）", "type": "购物站 Shopping stop", "city": "南投 Nantou", "stay_minutes": 60, "rebate_per_head": 200, "commission_pct": 0.1,
               "mandatory": "Y", "notes": "全程不進購物站每人加 TWD 1,500"}
        snote = "请列出标准购物站、人头费与「免购物」加价，以及常见自费项目价格，以便 WEBUY 做纯玩与常规两个版本。"
    sheet(wb, ctx, "7 购物与自费", "7 购物站与自费 Shopping stops & Optional tours", snote, s_cols, sex, spre, blank=10, lists=lists)

    # 8 季节与节庆
    e_cols = [
        col("event", "季节 / 节庆", "Season / event", "시즌 / 축제", 26, "pre"), col("webuy_window", "WEBUY 参考时间", "Reference window", "참고 시기", 22, "pre"),
        col("dates_2027", "2027 确切日期", "2027 exact dates", "2027 일정", 18), col("city", "影响城市", "Cities affected", "영향 도시", 16),
        col("hotel_surcharge_pct", "酒店加价 %", "Hotel surcharge %", "호텔 할증", 10, num="0%"), col("blackout", "是否不能接团", "Blackout", "접수 불가", 9, dv="yn"),
        col("recommend", "地接建议（必去 / 回避 / 替代）", "DMC advice", "추천 / 대체", 30), col("notes", "备注", "Notes", "비고", 24),
    ]
    if ctx.kr:
        S = json.load(open(os.path.join(kb, "seasons.json"), encoding="utf-8"))
        epre = [{"event": s["name_cn"], "webuy_window": s["best"]} for s in S]
        epre += [{"event": "春节 Seollal", "webuy_window": "2027-02-05 ~ 02-09（以官方公布为准）"}, {"event": "中秋 Chuseok", "webuy_window": "2027-09-14 ~ 09-16（以官方公布为准）"},
                 {"event": "中国国庆黄金周（济州 / 首尔酒店）", "webuy_window": "10 月 1 日 ~ 7 日"}, {"event": "日本黄金周", "webuy_window": "4 月 29 日 ~ 5 月 5 日"},
                 {"event": "新加坡学校假期", "webuy_window": "6 月、11 月中 ~ 12 月"}]
        eex = {"event": "镇海军港节樱花", "webuy_window": "3 月下旬–4 月上旬", "dates_2027": "2027-03-27 ~ 04-05", "city": "昌原（镇海）/ 釜山", "hotel_surcharge_pct": 0.3,
               "blackout": "N", "recommend": "大巴只能停外围，需改坐接驳车", "notes": ""}
    else:
        epre = [{"event": e, "webuy_window": w} for e, w, imp in TW.SEASONS]
        eex = {"event": "阿里山花季（櫻花）", "webuy_window": "3 月中 ~ 4 月中", "dates_2027": "2027-03-15 ~ 04-15", "city": "嘉義 阿里山", "hotel_surcharge_pct": 0.2,
               "blackout": "N", "recommend": "假日管制，建議平日上山", "notes": "需換接駁車"}
    sheet(wb, ctx, "8 季节节庆", "8 季节与节庆 Seasons & Events", "灰色为 WEBUY 掌握的季节窗口；请填 2027 确切日期、酒店加价与操作建议，可补充其他节庆。", e_cols, eex, epre, blank=10, lists=lists)

    # 9 行程报价（Quotes + Quote_Lines）
    fx_a, fx_b = lists["fx_range"]
    q_cols = [
        col("rfq_code", "询价编号", "RFQ code", "견적 요청 번호", 12, "pre"), col("itinerary", "线路骨架", "Itinerary skeleton", "일정 개요", 30, "pre"),
        col("days", "天数", "Days", "일수", 6, "pre"), col("hotel_level", "酒店档次", "Hotel level", "호텔 등급", 10, "pre"),
        col("date_from", "出发期 起", "Departure from", "출발 시작", 11, num="yyyy-mm-dd"), col("date_to", "出发期 止", "Departure to", "출발 종료", 11, num="yyyy-mm-dd"),
        col("season", "季节档", "Season", "시즌", 11, dv="season"), col("pax_tier", "人数档", "Pax tier", "인원 구간", 9, dv="tier"),
        col("currency", "币种", "Currency", "통화", 8, dv="currency"),
        col("price_twin", "成人 / 人（两人一房）", "Adult/pax twin", "성인 1인 (2인 1실)", 12, num="#,##0"),
        col("single_supp", "单房差", "Single supplement", "싱글 차액", 11, num="#,##0"), col("child_with_bed", "儿童占床", "Child w/ bed", "아동 베드 포함", 11, num="#,##0"),
        col("child_no_bed", "儿童不占床", "Child no bed", "아동 베드 미포함", 11, num="#,##0"), col("triple", "三人房 / 人", "Triple/pax", "트리플 1인", 11, num="#,##0"),
        col("foc", "FOC（如 15+1）", "FOC", "FOC", 9),
        col("price_twin_sgd", "SGD 参考价（自动）", "SGD ref. (auto)", "SGD 환산 (자동)", 12, "calc", num="#,##0",
            note="= 成人价 × Lists 表参考汇率；仅供比价，正式成本以 Lark FX_Rates 为准。"),
        col("meals_count", "含餐（B/L/D 次数）", "Meals incl. (B/L/D)", "식사 포함", 12), col("incl_tickets", "含门票", "Tickets incl.", "입장료 포함", 8, dv="yn"),
        col("incl_guide", "含导游", "Guide incl.", "가이드 포함", 8, dv="yn"), col("incl_tips", "含小费", "Tips incl.", "팁 포함", 8, dv="yn"),
        col("incl_domestic_transport", "含国内段机票 / 高铁", "Domestic flight/rail incl.", "국내선/KTX 포함" if ctx.kr else "", 12, dv="yn"),
        col("shopping_stops", "购物站数", "Shopping stops", "쇼핑 횟수", 8), col("notes", "备注（附加费 / 不含项）", "Notes", "비고", 28),
    ]
    if ctx.kr:
        F = {f["id"]: f for f in json.load(open(os.path.join(kb, "families.json"), encoding="utf-8"))}
        rfqs = []
        cn = {"Jeju": "济州", "Busan": "釜山", "Seoul": "首尔", "Daegu": "大邱", "Jeonju": "全州", "Gyeongju": "庆州", "Pyeongchang": "平昌"}
        for fid, days in (("K04", 8), ("K05", 8), ("K06", 8), ("K07", 8), ("K03", 8)):
            f = F[fid]
            rfqs.append((f"RFQ-{fid}", f["name_cn"] + "：" + " > ".join(f"{cn.get(s['city'], s['city'])} {s['nights']}N" for s in f["skeleton"]), days))
        rfqs.append(("RFQ-K01", "首尔 + 京畿深度：首尔 5N（每日放射）", 7))
    else:
        rfqs = [("RFQ-TW8", "台湾环岛：台北 2N > 花莲 1N > 台东 1N > 高雄 1N > 日月潭 1N > 台北 1N", 8),
                ("RFQ-TW5", "北台湾：台北 3N（九份 / 野柳 / 十分）> 宜兰礁溪 1N", 5),
                ("RFQ-TW6", "中南部：台中 1N > 阿里山 1N > 台南 1N > 高雄 1N > 台北 1N", 6),
                ("RFQ-TW6B", "台北 + 日月潭 + 台中：台北 2N > 日月潭 1N > 台中 2N", 6)]
    qpre = []
    for code, sk, days in rfqs:
        for tier in ("10–15", "16–20", "21–25", "26–30", "31–35"):
            qpre.append({"rfq_code": code, "itinerary": sk, "days": days, "hotel_level": "4★", "pax_tier": tier})
    qex = {"rfq_code": "RFQ-K04", "itinerary": "济州 2N > 釜山 2N > 首尔 2N", "days": 8, "hotel_level": "4★", "date_from": datetime.date(2027, 4, 15), "date_to": datetime.date(2027, 6, 30),
           "season": "平季 Shoulder", "pax_tier": "21–25", "currency": cur, "price_twin": 1150000 if ctx.kr else 26000, "single_supp": 380000 if ctx.kr else 8000,
           "child_with_bed": 1100000 if ctx.kr else 25000, "child_no_bed": 900000 if ctx.kr else 20000, "triple": 1120000 if ctx.kr else 25500, "foc": "15+1",
           "meals_count": "6B 6L 5D", "incl_tickets": "Y", "incl_guide": "Y", "incl_tips": "N", "incl_domestic_transport": "Y", "shopping_stops": 3, "notes": "4/1–4/10 樱花期每人 +KRW 50,000" if ctx.kr else "清明連假每人 +TWD 1,500"}
    if not ctx.kr: qex.update({"rfq_code": "RFQ-TW8", "itinerary": "台北 2N > 花蓮 1N > 台東 1N > 高雄 1N > 日月潭 1N > 台北 1N"})
    cidx = {c["key"]: get_column_letter(i) for i, c in enumerate(q_cols, 1)}
    fx = {"price_twin_sgd": lambda r: f'=IF(OR({cidx["price_twin"]}{r}="",{cidx["currency"]}{r}=""),"",{cidx["price_twin"]}{r}*INDEX({fx_b},MATCH({cidx["currency"]}{r},{fx_a},0)))'}
    ws = sheet(wb, ctx, "9 行程报价", "9 行程报价 Package Quote（按 出发期 × 人数档 一行）",
               "对应 Lark Quotes / Quote_Lines。灰色为 WEBUY 询价线路与人数档；请按出发期段填写每人净价（同一线路不同季节可复制行）。行程细节请填「10 逐日行程」。绿色列自动计算。",
               q_cols, qex, qpre, blank=20, formulas=fx, lists=lists)
    ws.cell(EX_ROW, q_cols.index(next(c for c in q_cols if c["key"] == "price_twin_sgd")) + 1).value = fx["price_twin_sgd"](EX_ROW)

    # 10 逐日行程
    d_cols = [
        col("rfq_code", "询价编号", "RFQ code", "견적 요청 번호", 12), col("day", "天", "Day", "일차", 6), col("route", "城市路线", "Route", "이동 경로", 22),
        col("transport", "交通（车程 / 航班）", "Transport", "교통", 18), col("attractions", "景点（用「3 景点门票」里的名称，顿号分隔）", "Attractions", "관광지", 40),
        col("breakfast", "早餐", "Breakfast", "조식", 14), col("lunch", "午餐", "Lunch", "중식", 18), col("dinner", "晚餐", "Dinner", "석식", 18),
        col("hotel", "住宿酒店", "Hotel", "숙박", 24), col("driving_hours", "当天车程（小时）", "Driving hours", "운전 시간", 9, num="0.0"),
        col("shopping", "购物站", "Shopping stop", "쇼핑", 14), col("notes", "备注（季节替代方案）", "Notes", "비고", 26),
    ]
    dex = {"rfq_code": "RFQ-K04", "day": 2, "route": "仁川 > 金浦 ✈ 济州", "transport": "大巴 1h + GMP–CJU 1h10m", "attractions": "爱妓峰和平生态公园、金浦米花农场、雪绿茶博物馆",
           "breakfast": "机上", "lunch": "石锅拌饭", "dinner": "韩定食", "hotel": "Jeju Best Western 或同级", "driving_hours": 3, "shopping": "", "notes": "12–3 月牛岛停航改城山日出峰"} if ctx.kr else \
          {"rfq_code": "RFQ-TW8", "day": 3, "route": "台北 > 宜蘭 > 花蓮", "transport": "遊覽車 3h", "attractions": "國立傳統藝術中心、清水斷崖、七星潭", "breakfast": "酒店", "lunch": "宜蘭風味餐",
           "dinner": "原住民風味餐", "hotel": "花蓮翰品酒店 或同級", "driving_hours": 3.5, "shopping": "", "notes": "蘇花改管制時改搭火車"}
    sheet(wb, ctx, "10 逐日行程", "10 逐日行程 Day-by-day Plan", "每个询价编号一组逐日行程，WEBUY 会用它生成双语行程 PDF 并跑合理性校验（景点归属、车程、季节）。", d_cols, dex, [], blank=60, lists=lists)

    # 填写说明
    R = readme
    R["A1"] = ctx.t(f"WEBUY TRAVEL · {country}地接资料与报价模板（{VERSION}）"); R["A1"].font = Font(name=FONT, size=16, bold=True, color="1F3864")
    lines = [
        ("用途", f"统一{country}地接社的资源与报价格式，填回后由 WEBUY 自动转入产品库（Lark Product Base：Suppliers / Quotes / Quote_Lines），用于比价、成本核算与行程生成。"),
        ("Purpose", f"One standard format for {'Korea' if ctx.kr else 'Taiwan'} ground handlers. Returned files are imported automatically into WEBUY's product base for costing and itinerary building."),
        ("颜色", "黄色 = 请填写 · 灰色 = WEBUY 预填（请勿修改，可在备注补充）· 绿色 = 公式自动计算 · 浅灰斜体第 5 行 = 示例，不会读入"),
        ("表头", "第 3 行为字段名（中文 / English" + (" / 한국어" if ctx.kr else "") + "），第 4 行灰色小字为系统字段 key，请勿删除或改动。"),
        ("币种", f"默认 {cur}；如用 USD / SGD 报价请在每行「币种」列注明。所有价格为净价（net），含当地税。"),
        ("日期", "日期格式 YYYY-MM-DD；季节档：淡季 / 平季 / 旺季 / 节假，与 WEBUY 出发日期表一致。"),
        ("新增行", "预填清单不够时直接在下方黄色空行继续填写；不够可插入行（保持列顺序）。"),
        ("不能操作", "预填的酒店 / 景点如不能操作或不建议，价格留空并在「备注」写原因或替代建议。"),
        ("交回", "填好后整份 Excel 交回 WEBUY Planner（请勿转成 PDF / 图片）。原始报价单可另附，但以本模板为准。"),
        ("表单", "1 供应商信息 · 2 酒店 · 3 景点门票 · 4 餐食 · 5 用车交通 · 6 导游领队 · 7 购物与自费 · 8 季节节庆 · 9 行程报价 · 10 逐日行程 · Lists（下拉选项与参考汇率）"),
    ]
    if ctx.kr:
        lines.insert(2, ("안내", "WEBUY TRAVEL(싱가포르) 표준 견적 양식입니다. 노란색 셀에 입력해 주시고, 회색 셀은 WEBUY가 미리 입력한 항목입니다(수정 금지). 금액은 KRW 넷 요금(세금 포함) 기준입니다."))
        lines.append(("数据来源", "灰色预填来自 WEBUY 韩国线知识库（korea/kb，2026-10-08 抓取新加坡 3 家旅行社 + 首尔地接在售线路 53 条），代表市场上实际在卖的酒店、景点与路段。"))
    else:
        lines.append(("数据来源", "灰色预填为 WEBUY 整理的台湾团体常用景点 / 路段 / 节庆清单（尚无竞品抓取数据），请地接按实际资源修正与补充。"))
    for i, (k, v) in enumerate(lines, 3):
        a = R.cell(i, 1, ctx.t(k)); a.font = Font(name=FONT, bold=True); a.alignment = WRAP
        b = R.cell(i, 2, ctx.t(v) if k not in ("Purpose", "안내") else v); b.font = F_BODY; b.alignment = WRAP
        R.row_dimensions[i].height = 44
    leg = len(lines) + 4
    for j, (fill, txt) in enumerate([(FILL_INPUT, "请填写 Fill in"), (FILL_PRE, "WEBUY 预填 Pre-filled"), (FILL_CALC, "自动计算 Formula"), (FILL_EX, "示例 Example")]):
        c = R.cell(leg, 1 + j * 0, None)
        c = R.cell(leg + j, 1, ctx.t(txt)); c.fill = fill; c.font = F_BODY; c.border = BORDER
    R.column_dimensions["A"].width = 22; R.column_dimensions["B"].width = 120
    wb.move_sheet("Lists", offset=len(wb.sheetnames))
    wb.active = 0
    wb.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--kb", default=os.path.join(HERE, "..", "korea", "kb")); ap.add_argument("--out", default=HERE)
    a = ap.parse_args()
    p1 = build(Ctx("KR"), a.kb, os.path.join(a.out, "WEBUY_DMC_Template_Korea.xlsx"))
    p2 = build(Ctx("TW"), a.kb, os.path.join(a.out, "WEBUY_DMC_Template_Taiwan.xlsx"))
    print(p1); print(p2)

if __name__ == "__main__":
    main()
